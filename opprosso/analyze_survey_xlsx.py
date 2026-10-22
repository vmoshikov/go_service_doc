#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Анализ Excel с результатами опроса (сырой экспорт или сводка Oprosso™).

Структура страницы (эвристика analyze_survey_xlsx):
  • Строка «N. Название» — начало блока вопроса.
  • Далее либо матрица (пустая/заголовочная строка, строки = сущности, столбцы = ответы),
    либо «Вариант ответа» + таблица вариантов.

Два формата данных (detect_data_format):
  raw — в столбцах B… лежат ответы отдельных респондентов;
  oprosso_summary — агрегаты: «Кол-во ответов», «Количество ответивших: N», матрица из чисел.

CLI: -q список вопросов, -r отчёт, --report-hints full|brief|off — подсказки в текстовом отчёте;
--exclude-unused-tools — убрать инструменты без использования по матрице частоты
(--tool-question, по умолчанию 9).
  --exclude-3 — исключать оценку «3» как «нет опыта»; по умолчанию оценка 3 учитывается.

Для страниц по инструменту (oprosso): по матрице Q9 считается число респондентов «без опыта»
(«не знаком», «знаю, но не использую», не ответили по строке); оно записывается в
tool_without_experience_n и вычитается из счётчика оценки «3» на каждой странице и в каждом
подвопросе по этому инструменту. Узкая сумма по колонкам «не использую»/… — в
tool_matrix_non_usage_n (справочно).
"""

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple, Union

try:
    import openpyxl
except ImportError:
    print("Требуется openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


# Паттерн: число. название (например "1. Ваша роль в команде")
PAGE_HEADER_PATTERN = re.compile(r"^\s*(\d+)\.\s+(.+)$")

# Страница «по инструменту»: "GigaChat .. …" или "13. GigaChat Работа с контекстом …"
DETAIL_PAGE_DOTDOT_PATTERN = re.compile(r"^([^.]+?)\s*\.\.\s*.+$")
DETAIL_PAGE_NUM_TOOL_PATTERN = re.compile(r"^\s*\d+\.\s*([^\s]+)")

# Ключевые фразы
VARIANT_OTVETA = "Вариант ответа"

# Для подсчёта «с опытом» в матрице (не для фильтра страниц по инструменту)
EXCLUDE_PHRASES = (
    "не знаю",
    "не использую",
    "не используется",
)
NO_EXPERIENCE_VALUES = ("3",)

# Первое слово после «N.» не считаем инструментом (обычные вопросы)
_TOOL_PAGE_FIRST_WORD_STOP = frozenset({
    "ваша", "ваш", "ваше", "какой", "какая", "какие", "какое", "оцените",
    "укажите", "выберите", "опишите", "насколько", "в", "как", "что", "есть",
    "при", "если", "для", "от", "до", "по", "на", "вопрос", "пройдите",
    "роль", "опыт", "уровень", "частота", "использование", "знакомство",
})

# Ответы в матрице инструментов → исключить респондента со страниц вопросов по этому инструменту
TOOL_DETAIL_EXCLUDE_PHRASES = (
    "не знаком",
    "знаю, но не использую",
    "знаю но не использую",
    "знаю, но не используется",
    "знаю но не используется",
)


def _category_is_non_usage_only(cat: str) -> bool:
    """
    Категория в матрице «частота/использование»: только «не знаком» / «знаю, но не использую».
    Ответы в этих колонках не считаются фактическим использованием инструмента.
    """
    a = cat.lower().strip()
    if "не знаком" in a:
        return True
    if "знаю" in a and "не использ" in a:
        return True
    return False


def _category_is_matrix_non_usage_subtract(cat: str) -> bool:
    """
    Матрица Q9: колонки, которые входят в «не использую» для вычитания из оценки «3»
    на страницах по инструменту — «не использую», «знаю, но не использую».
    «Не знаком» сюда не входит (отдельная колонка).
    """
    a = cat.lower().strip()
    if "не знаком" in a:
        return False
    if "знаю" in a and "не использ" in a:
        return True
    if "не использую" in a:
        return True
    return False


def _category_is_regular_or_often_use(cat: str) -> bool:
    """
    Явное использование: колонки «использую регулярно», «использую иногда».
    «Пробовал» сюда не входит — иначе почти все инструменты считаются «используемыми».
    """
    a = cat.lower().strip()
    if "регулярно" in a:
        return True
    if "использую" in a and "иногда" in a:
        return True
    return False


def _tool_matrix_row_has_usage(
    row_counts: Dict[str, int],
    policy: str = "regular_or_often",
) -> bool:
    """
    policy=regular_or_often (по умолчанию для --exclude-unused-tools):
        использование = есть ненулевой ответ в «регулярно» или «иногда».
    policy=any_beyond_non_usage:
        как раньше: любая колонка кроме «не знаком» / «знаю, но не использую»
        (в т.ч. «пробовал» считается использованием).
    """
    for cat, n in row_counts.items():
        if not n or n <= 0:
            continue
        if policy == "any_beyond_non_usage":
            if not _category_is_non_usage_only(cat):
                return True
        else:
            if _category_is_regular_or_often_use(cat):
                return True
    return False


def _get_cell_str(value) -> str:
    """Преобразует значение ячейки в строку."""
    if value is None:
        return ""
    return str(value).strip()


def _is_empty_row(row: tuple) -> bool:
    """Проверяет, пустая ли строка (все ячейки пустые или None)."""
    return all(_get_cell_str(c) == "" for c in row)


def _first_cell(row: tuple) -> str:
    """Первая ячейка строки."""
    return _get_cell_str(row[0]) if row else ""


def _has_variant_otveta(row: tuple) -> bool:
    """Есть ли в строке 'Вариант ответа'."""
    return any(VARIANT_OTVETA in _get_cell_str(c) for c in row)


def _is_page_header(cell: str) -> Tuple[bool, Optional[str], Optional[int]]:
    """
    Проверяет, является ли строка заголовком страницы (N. Название).
    Возвращает (True, название, номер) или (False, None, None).
    """
    m = PAGE_HEADER_PATTERN.match(cell)
    if m:
        return True, m.group(2).strip(), int(m.group(1))
    return False, None, None


def normalize_tool_key(name: str) -> str:
    """Ключ для сопоставления инструмента в матрице и в заголовке страницы."""
    return name.strip().lower()


def extract_tool_from_page_name(page_name: str) -> Optional[str]:
    """
    Имя инструмента для страниц вопросов по инструменту:
    "GigaChat .. …" или "13. GigaChat Работа с контекстом …"
    Если в имени нет префикса «N.» (как в поле name из analyze_survey_xlsx) —
    берётся первое слово, если оно не из стоп-листа обычных вопросов.
    """
    s = page_name.strip()
    m = DETAIL_PAGE_DOTDOT_PATTERN.match(s)
    if m:
        return m.group(1).strip()
    m2 = DETAIL_PAGE_NUM_TOOL_PATTERN.match(s)
    if m2:
        w = m2.group(1).strip()
        if w.lower() in _TOOL_PAGE_FIRST_WORD_STOP:
            return None
        return w
    parts = s.split(None, 1)
    if parts:
        w = parts[0].strip()
        if w and w.lower() not in _TOOL_PAGE_FIRST_WORD_STOP:
            return w
    return None


def _answer_should_exclude(answer: str, exclude_rating_3: bool = False) -> bool:
    """Проверяет, исключает ли ответ респондента (нет опыта с инструментом)."""
    a = answer.lower().strip()
    if any(phrase in a for phrase in EXCLUDE_PHRASES):
        return True
    if exclude_rating_3 and a in NO_EXPERIENCE_VALUES:
        return True
    return False


def _answer_excludes_from_tool_detail(answer: str) -> bool:
    """Ответ в матрице инструментов → не показывать этого респондента на страницах по инструменту."""
    a = answer.lower().strip()
    return any(phrase in a for phrase in TOOL_DETAIL_EXCLUDE_PHRASES)


def _answer_key_counts_as_non_usage_for_denom(key: str) -> bool:
    """
    Вариант ответа «не использую» / «знаю, но не использую» (и т.п.) —
    количество по таким строкам вычитается из базы tool_experience_n при % по странице инструмента.
    """
    a = key.lower().strip()
    if "не знаком" in a:
        return True
    if "знаю" in a and "не использ" in a:
        return True
    if "не использую" in a:
        return True
    if "не используется" in a and "использую" not in a:
        return True
    return False


def _non_usage_answer_count(answers: Dict[str, int]) -> int:
    """Сумма количеств по вариантам «не использую» / «знаю, но не использую» и т.д."""
    n = 0
    for k, v in (answers or {}).items():
        if v and _answer_key_counts_as_non_usage_for_denom(str(k)):
            n += int(v)
    return n


def _answers_are_pure_likert_1_5_keys(answers: Dict[str, int]) -> bool:
    """Варианты ответа — только «1»…«5» (одна цифра), без текста."""
    if not answers:
        return False
    for k in answers:
        ks = str(k).strip()
        if not (ks.isdigit() and len(ks) == 1):
            return False
        n = int(ks)
        if not 1 <= n <= 5:
            return False
    return True


# Парсинг «3 — текст» в _likert_score_from_key / _sort_answers_for_report
RATING_SORT_PATTERN = re.compile(r"^\s*(\d{1,2})\b")


def _likert_score_from_key(key: Any) -> Optional[int]:
    """
    Оценка 1–5 из подписи варианта в сводке Oprosso:
    «3», «3 — …», «5. Полностью», число float и т.д.
    """
    if key is None:
        return None
    ks = str(key).strip()
    if not ks:
        return None
    if ks.isdigit() and len(ks) == 1:
        n = int(ks)
        if 1 <= n <= 5:
            return n
    try:
        fv = float(ks.replace(",", "."))
        if abs(fv - round(fv)) < 1e-9:
            ri = int(round(fv))
            if 1 <= ri <= 5:
                return ri
    except ValueError:
        pass
    m = RATING_SORT_PATTERN.match(ks)
    if m:
        n = int(m.group(1))
        if 1 <= n <= 5:
            return n
    return None


def _find_rating_3_key(answers: Dict[str, int]) -> Optional[Any]:
    """Ключ варианта, соответствующий оценке 3 (после парсинга подписи)."""
    for k in answers:
        if _likert_score_from_key(k) == 3:
            return k
    return None


def _lookup_tool_matrix_map(m: Dict[str, int], page_tool: str) -> Optional[int]:
    """
    Значение по инструменту из матрицы: точное совпадение ключа или
    вхождение имени (cursor ↔ «Cursor IDE» в строке матрицы).
    """
    if not m:
        return None
    k = normalize_tool_key(page_tool)
    if k in m:
        return m[k]
    for mk, v in m.items():
        if len(k) >= 3 and len(mk) >= 3 and (k in mk or mk in k):
            return v
    # Первое слово совпадает: «Cursor …» ↔ «Cursor AI»
    toks = k.split()
    ptok = toks[0] if toks else k
    if ptok:
        for mk, v in m.items():
            mtoks = mk.split()
            mtok = mtoks[0] if mtoks else mk
            if ptok == mtok:
                return v
    return None


def _resolve_n_subtract_for_rating_3(s: Dict[str, Any]) -> Optional[int]:
    """
    Сколько вычесть из оценки «3» на странице/подвопросе по инструменту:
    1) tool_without_experience_n — полная сумма респондентов «без опыта» по матрице Q9;
    2) иначе tool_matrix_non_usage_n (узкая оценка по колонкам «не использую»/…);
    3) иначе max(0, всего − N с опытом), если известны total_respondents и tool_experience_n.
    """
    tr = int(s["total_respondents"] or 0)
    te_opt = s.get("tool_experience_n")
    te = int(te_opt) if te_opt is not None else None
    tw = s.get("tool_without_experience_n")
    if tw is not None:
        return int(tw)
    n_mu = s.get("tool_matrix_non_usage_n")
    if n_mu is not None:
        return int(n_mu)
    if te is not None and te > 0 and tr > te:
        return max(0, tr - te)
    return None


def _tool_page_instrument_answers_adjusted(
    s: Dict[str, Any],
    answers: Dict[str, int],
) -> Optional[Tuple[Dict[str, int], int]]:
    """
    Страница инструмента + сводка Oprosso: из счётчика варианта «3» вычитается
    tool_without_experience_n (полная сумма «без опыта» по матрице Q9 для инструмента),
    либо tool_matrix_non_usage_n / fallback по total_respondents − tool_experience_n.

    Одно и то же N применяется ко всем подвопросам на странице (multiple_questions).

    Проценты: от tool_experience_n (N с опытом), если сумма по вариантам не больше;
    иначе от суммы (переполнение сводки).
    """
    if s.get("data_format") != "oprosso_summary" or not s.get("tool"):
        return None
    if not answers:
        return None
    key3 = _find_rating_3_key(answers)
    if key3 is None:
        return None
    n_sub = _resolve_n_subtract_for_rating_3(s)
    if n_sub is None:
        return None
    te_opt = s.get("tool_experience_n")
    te = int(te_opt) if te_opt is not None else None
    out: Dict[str, int] = {}
    for k, v in answers.items():
        if k == key3:
            out[k] = max(0, int(v) - n_sub)
        else:
            out[k] = int(v)
    s_rest = sum(out.values())
    if te is not None:
        te = int(te)
        if te > 0:
            if s_rest > te:
                denom = max(1, s_rest)
            else:
                denom = max(1, te)
        else:
            denom = max(1, s_rest)
    else:
        denom = max(1, s_rest)
    return out, denom


def _sort_answers_for_report(
    items: List[Tuple[str, int]],
) -> List[Tuple[str, int]]:
    """
    Сортировка для отчёта: оценки шкалы 1–5 — по убыванию (5, 4, 3, 2, 1);
    прочие варианты — по убыванию частоты, затем по тексту.
    """

    def _key(it: Tuple[str, int]) -> Tuple[int, int, str]:
        label, cnt = it
        t = label.strip()
        if t.isdigit():
            n = int(t)
            if 1 <= n <= 5:
                return (0, -n, "")
        m = RATING_SORT_PATTERN.match(label)
        if m:
            n = int(m.group(1))
            if 1 <= n <= 5:
                return (0, -n, "")
        return (1, -cnt, t.lower())

    return sorted(items, key=_key)


def _parse_int_cell(val) -> Optional[int]:
    """Парсит число из ячейки (агрегированные количества)."""
    if val is None:
        return None
    s = str(val).strip().replace("\u00a0", "").replace(" ", "")
    if not s or "%" in s:
        return None
    try:
        return int(float(s.replace(",", ".")))
    except ValueError:
        return None


def detect_data_format(rows: List[tuple]) -> str:
    """
    oprosso_summary — сводка с «Кол-во ответов» / Oprosso™.
    raw — ответы по респондентам в столбцах.
    """
    for r in rows[:400]:
        if not r:
            continue
        for c in r[:12]:
            t = _get_cell_str(c).lower()
            if "oprosso" in t:
                return "oprosso_summary"
            if "кол-во ответов" in t or "количество ответов" in t:
                return "oprosso_summary"
    return "raw"


def get_main_questions(pages: List[dict]) -> List[Tuple[int, str]]:
    """Возвращает пронумерованный список основных вопросов (страниц)."""
    return [(p.get("num", p["row"]), p["name"]) for p in pages]


def load_survey_rows(filepath: Union[str, Path]) -> List[tuple]:
    """Загружает все строки листа."""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def build_respondent_exclusion(
    rows: List[tuple],
    pages: List[dict],
    tool_question_num: int = 9,
    total_cols: int = 0,
) -> Dict[str, Set[int]]:
    """
    Только для raw: обход матрицы вопроса tool_question_num; в каждой ячейке — текст ответа респондента.
    Строит множество столбцов для исключения по фразам TOOL_DETAIL_EXCLUDE_PHRASES.
    """
    exclusion: Dict[str, Set[int]] = {}
    # Ищем строки вопроса tool_question_num: страницы с этим номером и следующие
    # (матрица: инструменты по строкам, ответы по столбцам)
    tool_page_rows: List[int] = []
    for p in pages:
        if p.get("num") == tool_question_num:
            tool_page_rows.append(p["row"] - 1)  # 0-based
            break

    if not tool_page_rows:
        return exclusion

    tc = total_cols or (len(rows[0]) if rows else 0)
    scan_cols = set(range(1, max(tc, 1)))

    start_row = tool_page_rows[0]
    # Следующие строки до следующей страницы — строки инструментов
    next_page_row = len(rows)
    for p in pages:
        if p["row"] - 1 > start_row:
            next_page_row = p["row"] - 1
            break

    # Пропускаем строку заголовков (пустая или "Вариант ответа")
    data_start = start_row + 1
    if data_start < len(rows) and (
        _is_empty_row(rows[data_start]) or _has_variant_otveta(rows[data_start])
    ):
        data_start += 1

    for r in range(data_start, next_page_row):
        if r >= len(rows):
            break
        row = rows[r]
        tool_name = _get_cell_str(row[0]) if row else ""
        if not tool_name or _is_page_header(tool_name)[0]:
            continue
        for col_idx in scan_cols:
            if col_idx >= len(row):
                continue
            val = _get_cell_str(row[col_idx])
            if val and _answer_excludes_from_tool_detail(val):
                nk = normalize_tool_key(tool_name)
                exclusion.setdefault(nk, set()).add(col_idx)

    return exclusion


def get_valid_respondent_columns(
    exclude_for_tool: Optional[str],
    exclusion_map: Dict[str, Set[int]],
    base_columns: Set[int],
) -> Set[int]:
    """
    Возвращает индексы столбцов респондентов, которые учитываются.
    base_columns — допустимые столбцы для данной страницы (напр. B–F для вопроса 9).
    """
    valid = set(base_columns)
    if not exclude_for_tool:
        return valid
    k = normalize_tool_key(exclude_for_tool)
    if k in exclusion_map:
        valid -= exclusion_map[k]
    return valid


def _get_page_block_bounds(
    rows: List[tuple],
    pages: List[dict],
    page_idx: int,
) -> Tuple[int, int]:
    """Возвращает (start_row, end_row) 0-based для блока страницы."""
    start = pages[page_idx]["row"] - 1
    end = len(rows)
    if page_idx + 1 < len(pages):
        end = pages[page_idx + 1]["row"] - 1
    return start, end


def detect_layout(rows: List[tuple]) -> str:
    """
    Определяет раскладку: 'columns' = респонденты в столбцах (col 1..N),
    'rows' = респонденты в строках (row 1..N).
    Эвристика: если в первой строке есть паттерн N. в нескольких ячейках — заголовки по колонкам.
    Если в первом столбце много N. — структура по строкам.
    """
    if not rows:
        return "columns"
    first_row = rows[0]
    first_col_headers = sum(
        1 for c in first_row if c and PAGE_HEADER_PATTERN.match(_get_cell_str(c))
    )
    first_col_values = sum(
        1 for r in rows[:50] if r and PAGE_HEADER_PATTERN.match(_get_cell_str(r[0]))
    )
    if first_col_values > first_col_headers:
        return "columns"  # N. в первом столбце — классическая структура
    return "rows"  # N. в первой строке — респонденты в строках


def _find_respondent_count_in_range(rows: List[tuple], start: int, end: int) -> Optional[int]:
    """Ищет «Количество ответивших: N» в блоке."""
    for r in range(start, min(end, len(rows))):
        row = rows[r]
        if not row:
            continue
        a = _get_cell_str(row[0])
        if "Количество ответивших" in a:
            n = _parse_int_cell(row[1] if len(row) > 1 else None)
            if n is not None:
                return n
    return None


def _row_is_variant_header(row: tuple) -> bool:
    if not row or len(row) < 2:
        return False
    return (
        VARIANT_OTVETA in _get_cell_str(row[0])
        and "кол-во ответов" in _get_cell_str(row[1]).lower()
    )


def _is_oprosso_rating_value_row(row: tuple) -> bool:
    """
    Строка шкалы 1–5 с количеством в колонке B (сводка Oprosso).
    Такие строки не считаем заголовком подвопроса.
    """
    if not row or len(row) < 2:
        return False
    x0 = _get_cell_str(row[0]).strip()
    if x0 not in ("1", "2", "3", "4", "5"):
        return False
    if _parse_int_cell(row[1]) is None:
        return False
    return True


def _is_oprosso_subquestion_label_row(row: tuple) -> bool:
    """
    Строка с текстом подвопроса перед блоком «Вариант ответа».
    В Oprosso после «Среднее значение» идёт следующая строка с названием подпункта
    (она попадает во внутренний цикл сбора ответов и раньше терялась).
    """
    if not row:
        return False
    x0 = _get_cell_str(row[0])
    if not x0:
        return False
    if _is_page_header(x0)[0]:
        return False
    if VARIANT_OTVETA in x0:
        return False
    if "среднее" in x0.lower() or "средн" in x0.lower():
        return False
    if _is_oprosso_rating_value_row(row):
        return False
    return True


def _filter_oprosso_answers_exclude_rating(
    ans: Dict[str, int],
    exclude_rating_3: bool,
) -> Dict[str, int]:
    """В сводке Oprosso убирает вариант «3» (нет опыта), если включено exclude_rating_3."""
    if not exclude_rating_3:
        return dict(ans)
    return {
        k: v
        for k, v in ans.items()
        if str(k).strip() not in NO_EXPERIENCE_VALUES
    }


def extract_oprosso_question_stats(
    rows: List[tuple],
    pages: List[dict],
    page_idx: int,
    exclude_rating_3: bool = False,
) -> Dict[str, Any]:
    """
    Парсинг сводки Oprosso: колонка «Кол-во ответов», матрица с числами.
    total_respondents берётся из «Количество ответивших: N».
    """
    p = pages[page_idx]
    start_row, end_row = _get_page_block_bounds(rows, pages, page_idx)
    page_type = p["type"]

    result: Dict[str, Any] = {
        "page_num": p["num"],
        "page_name": p["name"],
        "type": page_type if page_type != "unknown" else "single_question",
        "tool": p.get("tool"),
        "total_respondents": 0,
        "answers": {},
        "matrix": {},
        "matrix_n": {},
        "subquestions": [],
        "data_format": "oprosso_summary",
    }

    n_resp = _find_respondent_count_in_range(rows, start_row, end_row)
    result["total_respondents"] = n_resp or 0

    if start_row >= len(rows):
        return result

    # Матрица Oprosso: строка с пустым A — заголовки категорий в B, C, …; далее строки «инструмент | числа».
    if page_type == "matrix":
        header_r = None
        categories: List[str] = []
        for r in range(start_row + 1, end_row):
            row = rows[r]
            if not row or len(row) < 2:
                continue
            a0 = _get_cell_str(row[0])
            if "Количество ответивших" in a0:
                break
            if not a0:
                # строка заголовков категорий
                cats = []
                for j in range(1, len(row)):
                    t = _get_cell_str(row[j])
                    if t and _parse_int_cell(t) is None:
                        cats.append(t)
                if len(cats) >= 2:
                    header_r = r
                    categories = cats
                    break
        if header_r is not None and categories:
            mat: Dict[str, Dict[str, int]] = {}
            mat_n: Dict[str, int] = {}
            for r in range(header_r + 1, end_row):
                row = rows[r]
                if not row:
                    continue
                tool = _get_cell_str(row[0])
                if not tool or "Количество ответивших" in tool:
                    break
                if _is_page_header(tool)[0]:
                    break
                row_d: Dict[str, int] = {}
                acc_exp = 0
                for j, cat in enumerate(categories):
                    col_idx = j + 1
                    if col_idx >= len(row):
                        break
                    v = _parse_int_cell(row[col_idx])
                    if v is not None:
                        row_d[cat] = v
                        if not _category_is_non_usage_only(cat):
                            acc_exp += v
                mat[tool] = row_d
                # База для %: только «с опытом» (без «не знаком» / «знаю, но не использую»)
                mat_n[tool] = acc_exp
            result["matrix"] = mat
            result["matrix_n"] = mat_n
        return result

    # Один или несколько блоков «Вариант ответа | Кол-во ответов»; перед блоком может быть название подвопроса.
    subqs: List[Dict[str, Any]] = []
    r = start_row + 1
    current_subq = ""
    while r < end_row:
        row = rows[r]
        if not row:
            r += 1
            continue
        a0 = _get_cell_str(row[0])
        if "Количество ответивших" in a0:
            break
        if _is_page_header(a0)[0] and r > start_row + 1:
            break
        if _row_is_variant_header(row):
            sub_ans: Dict[str, int] = defaultdict(int)
            # Заголовок подвопроса задаётся строками перед «Вариант ответа» (current_subq)
            # или строкой после «Среднее» внутри предыдущего блока (pending_subq).
            block_title = current_subq or p["name"]
            current_subq = ""
            pending_subq = ""
            r += 1
            while r < end_row:
                row2 = rows[r]
                if not row2:
                    r += 1
                    continue
                x0 = _get_cell_str(row2[0])
                if _row_is_variant_header(row2):
                    break
                if "Количество ответивших" in x0:
                    break
                if _is_page_header(x0)[0]:
                    break
                if "среднее" in x0.lower() or "средн" in x0.lower():
                    r += 1
                    continue
                if _is_oprosso_subquestion_label_row(row2):
                    pending_subq = x0
                    r += 1
                    continue
                cnt = _parse_int_cell(row2[1] if len(row2) > 1 else None)
                if cnt is not None and x0:
                    sub_ans[x0] += cnt
                elif x0 and not cnt:
                    pass
                r += 1
            title = block_title
            subqs.append({
                "question": title,
                "answers": _filter_oprosso_answers_exclude_rating(
                    dict(sub_ans), exclude_rating_3
                ),
            })
            current_subq = pending_subq
            continue
        if a0 and not _row_is_variant_header(row) and "Количество ответивших" not in a0:
            if not any(_parse_int_cell(row[j] if j < len(row) else None) for j in range(1, min(4, len(row)))):
                current_subq = a0
        r += 1

    result["subquestions"] = subqs
    if len(subqs) == 1:
        result["answers"] = subqs[0]["answers"]
        result["type"] = "single_question"
    elif subqs:
        merged: Dict[str, int] = defaultdict(int)
        for sq in subqs:
            for k, v in sq["answers"].items():
                merged[f"{sq['question'][:40]}::{k}"] = v
        result["answers"] = dict(merged)
        result["type"] = "multiple_questions"
    return result


def _oprosso_matrix_experience_sum(row_d: Dict[str, int]) -> int:
    """Сумма по колонкам матрицы, считающимся «с опытом» (без «не знаком» / «знаю, но не использую»)."""
    n = 0
    for cat, v in row_d.items():
        if v and v > 0 and not _category_is_non_usage_only(cat):
            n += v
    return n


def _tool_row_without_experience_n(row_d: Dict[str, int], survey_n: int) -> int:
    """
    Число респондентов «без опыта» по строке матрицы частоты (Q9 и т.п.):
    сумма по колонкам «не знаком», «знаю, но не использую» (см. _category_is_non_usage_only)
    плюс не ответившие по строке (разница до «Количество ответивших»).
    Эту величину вычитаем из варианта «3» на каждой странице/подвопросе по инструменту.
    """
    n_no_exp = 0
    for cat, v in row_d.items():
        if v and _category_is_non_usage_only(cat):
            n_no_exp += int(v)
    sum_row = sum(int(x) for x in row_d.values() if x is not None)
    gap = max(0, int(survey_n) - sum_row) if survey_n else 0
    return n_no_exp + gap


def compute_tool_without_experience_counts_oprosso(
    rows: List[tuple],
    pages: List[dict],
    tool_question_num: int = 9,
) -> Dict[str, int]:
    """
    По матрице вопроса tool_question_num: для каждого инструмента — сколько респондентов
    «без опыта» (не в базе «с опытом» по матрице). Используется для вычитания из ответа «3»
    на всех вопросах по этому инструменту (в т.ч. каждого подпункта на странице).
    """
    for i, p in enumerate(pages):
        if p.get("num") != tool_question_num:
            continue
        if p.get("type") != "matrix":
            continue
        st = extract_oprosso_question_stats(
            rows, pages, i, exclude_rating_3=False
        )
        mat = st.get("matrix") or {}
        survey_n = int(st.get("total_respondents") or 0)
        out: Dict[str, int] = {}
        for tool, row_d in mat.items():
            out[normalize_tool_key(tool)] = _tool_row_without_experience_n(
                row_d, survey_n
            )
        return out
    return {}


def compute_tool_matrix_non_usage_subtract_oprosso(
    rows: List[tuple],
    pages: List[dict],
    tool_question_num: int = 9,
) -> Dict[str, int]:
    """
    По матрице вопроса tool_question_num (частота): для каждого инструмента — узкая сумма
    колонок «не использую» / «знаю, но не использую» (без «не знаком») + не ответили.
    Для вычитания из «3» на страницах по инструменту используется полная сумма
    «без опыта» — см. compute_tool_without_experience_counts_oprosso и tool_without_experience_n.
    """
    for i, p in enumerate(pages):
        if p.get("num") != tool_question_num:
            continue
        if p.get("type") != "matrix":
            continue
        st = extract_oprosso_question_stats(
            rows, pages, i, exclude_rating_3=False
        )
        mat = st.get("matrix") or {}
        survey_n = int(st.get("total_respondents") or 0)
        out: Dict[str, int] = {}
        for tool, row_d in mat.items():
            nu = 0
            for cat, v in row_d.items():
                if v and _category_is_matrix_non_usage_subtract(cat):
                    nu += int(v)
            sum_row = sum(int(x) for x in row_d.values() if x is not None)
            if survey_n > 0:
                nu += max(0, survey_n - sum_row)
            out[normalize_tool_key(tool)] = nu
        return out
    return {}


def compute_tool_experience_counts_oprosso(
    rows: List[tuple],
    pages: List[dict],
    tool_question_num: int = 9,
) -> Dict[str, int]:
    """
    По матрице вопроса tool_question_num (частота использования инструментов):
    для каждого инструмента — число респондентов «с опытом» (сумма колонок кроме
    «не знаком» и «знаю, но не использую»). Используется как база N для страниц
    вопросов по этому инструменту в сводке Oprosso.
    """
    for i, p in enumerate(pages):
        if p.get("num") != tool_question_num:
            continue
        if p.get("type") != "matrix":
            continue
        st = extract_oprosso_question_stats(
            rows, pages, i, exclude_rating_3=False
        )
        mat = st.get("matrix") or {}
        out: Dict[str, int] = {}
        for tool, row_d in mat.items():
            out[normalize_tool_key(tool)] = _oprosso_matrix_experience_sum(row_d)
        return out
    return {}


def enrich_oprosso_tool_pages_with_experience_n(
    stats_list: List[Dict[str, Any]],
    rows: List[tuple],
    pages: List[dict],
    tool_question_num: int = 9,
) -> List[Dict[str, Any]]:
    """
    Для страниц с полем tool добавляет:
    • tool_experience_n — N «с опытом» из матрицы Q9 (без «не знаком» / «знаю, но не использую»);
    • tool_without_experience_n — сумма респондентов «без опыта» по Q9 для инструмента
      («не знаком» + «знаю, но не использую» + не ответили по строке); из варианта «3»
      на каждой странице/подвопросе по этому инструменту вычитается это число;
    • tool_matrix_non_usage_n — узкая сумма по колонкам «не использую»/«знаю…» (без «не знаком»)
      + не ответили; сохраняется для справки (приоритет вычитания — tool_without_experience_n).

    total_respondents не меняется: общее число ответивших на опрос (из сводки).
    """
    exp = compute_tool_experience_counts_oprosso(rows, pages, tool_question_num)
    wo_map = compute_tool_without_experience_counts_oprosso(
        rows, pages, tool_question_num
    )
    nu_map = compute_tool_matrix_non_usage_subtract_oprosso(
        rows, pages, tool_question_num
    )
    if not exp and not wo_map and not nu_map:
        return stats_list
    out: List[Dict[str, Any]] = []
    for s in stats_list:
        t = s.get("tool")
        if not t:
            out.append(s)
            continue
        s2 = None
        n_exp = _lookup_tool_matrix_map(exp, t) if exp else None
        n_wo = _lookup_tool_matrix_map(wo_map, t) if wo_map else None
        n_nu = _lookup_tool_matrix_map(nu_map, t) if nu_map else None
        if n_exp is not None and n_exp > 0:
            if s2 is None:
                s2 = dict(s)
            s2["tool_experience_n"] = n_exp
        if n_wo is not None:
            if s2 is None:
                s2 = dict(s)
            s2["tool_without_experience_n"] = int(n_wo)
        if n_nu is not None:
            if s2 is None:
                s2 = dict(s)
            s2["tool_matrix_non_usage_n"] = int(n_nu)
        if s2 is not None:
            out.append(s2)
        else:
            out.append(s)
    return out


def _denom_oprosso_tool_question(
    s: Dict[str, Any],
    sum_ans: int,
    answers: Dict[str, Any],
) -> int:
    """
    Знаменатель % для вопроса на странице инструмента (oprosso).
    Из tool_experience_n вычитается сумма по вариантам «не использую» / «знаю, но не использую».
    Если сумма ответов по вариантам больше базы матрицы — % от суммы вариантов.
    """
    is_oprosso_tool = (
        s.get("data_format") == "oprosso_summary" and s.get("tool")
    )
    if not is_oprosso_tool:
        return s["total_respondents"] or sum_ans or 1
    te_raw = s.get("tool_experience_n")
    if te_raw is None:
        return s["total_respondents"] or sum_ans or 1
    te_raw = int(te_raw)
    if sum_ans > te_raw:
        return sum_ans
    nu = _non_usage_answer_count(answers or {})
    return max(1, te_raw - nu)


def compute_unused_tool_keys_oprosso(
    rows: List[tuple],
    pages: List[dict],
    tool_question_num: int = 9,
    policy: str = "regular_or_often",
) -> Set[str]:
    """
    Для сводки Oprosso: по матрице вопроса tool_question_num (обычно «Частота использования»)
    находит инструменты без использования по выбранной политике (см. --unused-tools-policy).
    """
    for i, p in enumerate(pages):
        if p.get("num") != tool_question_num:
            continue
        if p.get("type") != "matrix":
            continue
        st = extract_oprosso_question_stats(rows, pages, i)
        mat = st.get("matrix") or {}
        unused: Set[str] = set()
        for tool, row_d in mat.items():
            if not _tool_matrix_row_has_usage(row_d, policy=policy):
                unused.add(normalize_tool_key(tool))
        return unused
    return set()


def apply_exclude_unused_tools(
    stats_list: List[Dict[str, Any]],
    unused_keys: Set[str],
) -> List[Dict[str, Any]]:
    """
    Убирает из статистики строки матрицы по неиспользуемым инструментам и целые страницы
    вопросов, привязанные к такому инструменту (поле tool).
    """
    if not unused_keys:
        return stats_list
    out: List[Dict[str, Any]] = []
    for s in stats_list:
        tool_page = s.get("tool")
        if tool_page and normalize_tool_key(tool_page) in unused_keys:
            continue
        s2 = dict(s)
        mat = s2.get("matrix")
        if mat:
            s2["matrix"] = {
                k: v
                for k, v in mat.items()
                if normalize_tool_key(k) not in unused_keys
            }
            mn = s2.get("matrix_n") or {}
            s2["matrix_n"] = {
                k: v
                for k, v in mn.items()
                if normalize_tool_key(k) not in unused_keys
            }
        out.append(s2)
    return out


def extract_question_stats(
    rows: List[tuple],
    pages: List[dict],
    page_idx: int,
    valid_columns: Optional[Set[int]] = None,
    exclude_rating_3: bool = False,
    data_format: str = "raw",
) -> Dict[str, Any]:
    """
    Извлекает числовую статистику по ответам для одной страницы.
    valid_columns: индексы столбцов респондентов для учёта (None = все).
    data_format: raw — респонденты в столбцах; oprosso_summary — сводка с количествами.
    """
    if data_format == "oprosso_summary":
        return extract_oprosso_question_stats(
            rows, pages, page_idx, exclude_rating_3=exclude_rating_3
        )

    total_cols = len(rows[0]) if rows else 0
    p = pages[page_idx]
    start_row, end_row = _get_page_block_bounds(rows, pages, page_idx)
    page_type = p["type"]

    if valid_columns is None:
        valid_columns = set(range(1, total_cols))

    result = {
        "page_num": p["num"],
        "page_name": p["name"],
        "type": page_type,
        "tool": p.get("tool"),
        "total_respondents": len(valid_columns),
        "answers": defaultdict(int),
        "matrix": {},
        "matrix_n": {},  # по инструменту: кол-во респондентов с опытом (исключая "не использую")
        "subquestions": [],
    }

    if start_row >= len(rows):
        return dict(result)

    # Пропуск строки заголовка страницы
    data_start = start_row + 1
    if data_start >= end_row:
        return dict(result)

    header_row = rows[data_start]
    data_start += 1

    if page_type == "single_question":
        # Один вопрос: каждый респондент (колонка) даёт один ответ.
        # Сумма по колонкам: для каждой колонки берём первый непустой ответ в блоке.
        for col in valid_columns:
            val = None
            for r in range(data_start, end_row):
                if r >= len(rows):
                    break
                row = rows[r]
                if _has_variant_otveta(row) or _is_empty_row(row):
                    continue
                if col < len(row):
                    v = _get_cell_str(row[col])
                    if v and v != VARIANT_OTVETA and not _is_page_header(v)[0]:
                        val = v
                        break
            if val:
                result["answers"][val] += 1

    elif page_type == "matrix":
        # Строки = инструменты, столбцы = респонденты, ячейка = ответ.
        # Исключаем респондентов, ответивших "не использую" / "3" для данного инструмента —
        # считаем только тех, кто имеет опыт использования.
        if data_start < len(rows) and (
            _is_empty_row(rows[data_start]) or _has_variant_otveta(rows[data_start])
        ):
            data_start += 1
        for r in range(data_start, end_row):
            if r >= len(rows):
                break
            row = rows[r]
            row_label = _get_cell_str(row[0]) if row else ""
            if not row_label or _is_page_header(row_label)[0]:
                continue
            result["matrix"][row_label] = defaultdict(int)
            respondents_with_experience = 0
            for col in valid_columns:
                if col < len(row):
                    val = _get_cell_str(row[col])
                    if val:
                        if _answer_should_exclude(val, exclude_rating_3):
                            continue  # Исключаем: нет опыта с этим инструментом
                        result["matrix"][row_label][val] += 1
                        respondents_with_experience += 1
            result["matrix_n"][row_label] = respondents_with_experience

    elif page_type == "multiple_questions":
        # Строки = подвопросы (вопрос в col 0), столбцы = респонденты
        r = data_start
        while r < end_row and r < len(rows):
            row = rows[r]
            q_text = _get_cell_str(row[0]) if row else ""
            if not q_text or _is_page_header(q_text)[0]:
                r += 1
                continue
            sub_answers = defaultdict(int)
            for col in valid_columns:
                if col < len(row):
                    val = _get_cell_str(row[col])
                    if val:
                        sub_answers[val] += 1
            result["subquestions"].append({
                "question": q_text,
                "answers": dict(sub_answers),
            })
            r += 1

    result["answers"] = dict(result["answers"])
    result["matrix"] = {k: dict(v) for k, v in result["matrix"].items()}
    return result


def _get_page_column_bounds(rows: List[tuple]) -> List[Tuple[int, int, dict]]:
    """
    Для layout=rows: сканирует строку 0, находит заголовки N. и возвращает
    [(start_col, end_col, page_info), ...].
    """
    if not rows:
        return []
    header = rows[0]
    result = []
    i = 0
    while i < len(header):
        cell = _get_cell_str(header[i])
        is_h, name, num = _is_page_header(cell)
        if is_h and name:
            start = i
            end = i + 1
            j = i + 1
            while j < len(header):
                next_cell = _get_cell_str(header[j])
                if _is_page_header(next_cell)[0]:
                    end = j
                    break
                j += 1
            else:
                end = len(header)
            result.append((start, end, {"num": num, "name": name, "row": 1}))
            i = end
        else:
            i += 1
    return result


def extract_question_stats_rows_layout(
    rows: List[tuple],
    col_start: int,
    col_end: int,
    page_info: dict,
    valid_rows: Optional[Set[int]] = None,
) -> Dict[str, Any]:
    """
    Извлечение статистики когда респонденты в строках (rows 1..N).
    col_start, col_end — диапазон колонок для этого вопроса.
    """
    total_rows = len(rows) - 1 if rows else 0  # минус заголовок
    if valid_rows is None:
        valid_rows = set(range(1, len(rows)))
    result = {
        "page_num": page_info.get("num", 0),
        "page_name": page_info.get("name", ""),
        "type": "single_question",
        "tool": extract_tool_from_page_name(page_info.get("name", "")),
        "total_respondents": len(valid_rows),
        "answers": defaultdict(int),
        "matrix": {},
        "subquestions": [],
    }
    for r in valid_rows:
        if r >= len(rows):
            continue
        row = rows[r]
        for c in range(col_start, min(col_end, len(row))):
            val = _get_cell_str(row[c])
            if val:
                result["answers"][val] += 1
    result["answers"] = dict(result["answers"])
    return result


def extract_all_stats(
    rows: List[tuple],
    pages: List[dict],
    exclusion_map: Optional[Dict[str, Set[int]]] = None,
    tool_question_num: int = 9,
    layout: str = "columns",
    exclude_rating_3: bool = False,
    data_format: str = "raw",
) -> List[Dict[str, Any]]:
    """
    Собирает статистику по каждой странице из pages.
    oprosso_summary → только extract_oprosso_question_stats; raw → обход по столбцам респондентов.
    """
    if data_format == "oprosso_summary":
        layout = "columns"

    if layout == "rows":
        col_bounds = _get_page_column_bounds(rows)
        all_stats = []
        for start, end, page_info in col_bounds:
            stats = extract_question_stats_rows_layout(
                rows, start, end, page_info, None
            )
            all_stats.append(stats)
        return all_stats

    total_cols = len(rows[0]) if rows else 0
    all_stats = []

    for i, p in enumerate(pages):
        base_cols = set(range(1, total_cols))
        valid = None
        if (
            data_format == "raw"
            and exclusion_map
            and p.get("tool")
            and total_cols > 1
        ):
            valid = get_valid_respondent_columns(
                p["tool"], exclusion_map, base_cols
            )
        stats = extract_question_stats(
            rows,
            pages,
            i,
            valid,
            exclude_rating_3=exclude_rating_3,
            data_format=data_format,
        )
        all_stats.append(stats)

    if data_format == "oprosso_summary":
        all_stats = enrich_oprosso_tool_pages_with_experience_n(
            all_stats, rows, pages, tool_question_num
        )

    return all_stats


def _tool_scale_explanation(
    s: Dict[str, Any],
    tr: int,
    te_raw: Any,
    brief: bool,
    same_n_all_subquestions: bool = False,
) -> str:
    """Текст пояснения к вычитанию из «3» для full vs brief."""
    if s.get("tool_without_experience_n") is not None:
        if brief:
            if same_n_all_subquestions:
                return (
                    "N из матрицы частоты: «без опыта» по строке инструмента; "
                    "одинаково для всех подпунктов на странице"
                )
            return (
                "N из матрицы частоты: число респондентов «без опыта» по строке инструмента"
            )
        if same_n_all_subquestions:
            return (
                "респонденты «без опыта» по Q9 (не знаком + знаю, но не использую + "
                "не ответили по строке); то же N для всех подвопросов на странице"
            )
        return (
            "респонденты «без опыта» по Q9 для инструмента "
            "(«не знаком» + «знаю, но не использую» + не ответили по строке)"
        )
    if s.get("tool_matrix_non_usage_n") is not None:
        if brief:
            return "узкая сумма по матрице (колонки «не использую»/… + пропуски строки)"
        return (
            "узкая сумма по матрице Q9 (колонки «не использую»/«знаю, но не использую» "
            "+ не ответили по строке)"
        )
    te_s = te_raw if te_raw is not None else "—"
    if brief:
        return f"оценка по данным опроса: всего {tr}, с опытом по матрице {te_s}"
    return (
        f"оценка: всего {tr} − с опытом по матрице {te_s} "
        f"(строка инструмента в Q9 не сопоставлена — см. --tool-question)"
    )


def build_text_report(
    stats_list: List[Dict[str, Any]],
    validate: bool = False,
    report_hints: str = "full",
) -> str:
    """
    Формирует текстовый отчёт.

    report_hints:
      full  — пояснения по умолчанию (длинные);
      brief — короткий контекст для анализа (те же числа, меньше текста);
      off   — без подсказок в квадратных скобках и без расшифровки матрицы в шапке.
    """
    rh = report_hints if report_hints in ("full", "brief", "off") else "full"
    show_bracket = rh != "off"
    show_tool_meta = rh != "off"
    brief = rh == "brief"

    lines = []
    lines.append("=" * 80)
    lines.append("ОТЧЁТ ПО РЕЗУЛЬТАТАМ ОПРОСА")
    lines.append(f"Всего вопросов/страниц: {len(stats_list)}")
    total_resp = max((s["total_respondents"] for s in stats_list), default=0)
    lines.append(f"Макс. респондентов на странице: {total_resp}")
    fmt0 = stats_list[0].get("data_format") if stats_list else None
    if fmt0:
        lines.append(f"Формат: {fmt0}")
    lines.append("=" * 80)

    for s in stats_list:
        lines.append("")
        lines.append(f"{s['page_num']}. {s['page_name']}")
        if s.get("tool"):
            te = s.get("tool_experience_n")
            two = s.get("tool_without_experience_n")
            tnu = s.get("tool_matrix_non_usage_n")
            if show_tool_meta:
                parts = []
                if brief:
                    if te is not None:
                        parts.append(f"с опытом (матрица): {te}")
                    if two is not None:
                        parts.append(f"без опыта → вычитание из «3»: {two}")
                    if tnu is not None:
                        parts.append(f"узко «не исп.»: {tnu}")
                else:
                    if te is not None:
                        parts.append(f"База «с опытом» (матрица частоты): {te}")
                    if two is not None:
                        parts.append(
                            f"«без опыта» по матрице (вычитание из «3»): {two}"
                        )
                    if tnu is not None:
                        parts.append(
                            f"узкая сумма «не использую» по матрице (справочно): {tnu}"
                        )
                te_note = " | " + " | ".join(parts) if parts else ""
                lines.append(
                    f"   [Инструмент: {s['tool']}] | Респондентов: {s['total_respondents']}"
                    f"{te_note}"
                )
            else:
                lines.append(
                    f"   Инструмент: {s['tool']} | Респондентов: {s['total_respondents']}"
                )
        else:
            lines.append(f"   Респондентов: {s['total_respondents']}")
        lines.append("-" * 60)

        if s["type"] == "single_question" and s["answers"]:
            total_ans = sum(s["answers"].values())
            te_raw = s.get("tool_experience_n")
            is_oprosso_tool = (
                s.get("data_format") == "oprosso_summary"
                and s.get("tool")
                and total_ans > 0
            )
            exp_only = _tool_page_instrument_answers_adjusted(s, s["answers"])
            if exp_only is not None:
                filt_ans, denom = exp_only
                n_sub = _resolve_n_subtract_for_rating_3(s)
                tr = int(s["total_respondents"] or 0)
                nu_note = _tool_scale_explanation(s, tr, te_raw, brief, False)
                if show_bracket:
                    if brief:
                        lines.append(
                            f"   [1–5: из «3» вычтено N={n_sub}; знаменатель %: {denom}. "
                            f"{nu_note}]"
                        )
                    else:
                        lines.append(
                            f"   [Шкала 1–5: из счётчика «3» вычитается N={n_sub} — {nu_note}; "
                            f"знаменатель для %: {denom}]"
                        )
                if validate or (
                    show_bracket
                    and (total_ans != s["total_respondents"] and total_ans > 0)
                ):
                    lines.append(
                        f"   [Сумма ответов по колонкам: {total_ans}, респондентов: {s['total_respondents']}]"
                    )
                for ans, cnt in _sort_answers_for_report(list(filt_ans.items())):
                    pct = 100 * cnt / denom if denom else 0
                    lines.append(f"   {ans}: {cnt} ({pct:.1f}%)")
            else:
                if (
                    show_bracket
                    and is_oprosso_tool
                    and te_raw is not None
                    and total_ans > int(te_raw)
                ):
                    if brief:
                        lines.append(
                            f"   [Сумма вариантов {total_ans} > база «с опытом» {te_raw} — % от суммы вариантов]"
                        )
                    else:
                        lines.append(
                            f"   [Сводка: сумма по вариантам {total_ans}; база «с опытом» по матрице "
                            f"частоты инструментов N={te_raw} — проценты ниже от суммы ответов по вариантам]"
                        )
                if validate or (
                    show_bracket
                    and (total_ans != s["total_respondents"] and total_ans > 0)
                ):
                    lines.append(
                        f"   [Сумма ответов по колонкам: {total_ans}, респондентов: {s['total_respondents']}]"
                    )
                denom = _denom_oprosso_tool_question(s, total_ans, s["answers"])
                for ans, cnt in _sort_answers_for_report(list(s["answers"].items())):
                    pct = 100 * cnt / denom if denom else 0
                    lines.append(f"   {ans}: {cnt} ({pct:.1f}%)")

        elif s["type"] == "matrix" and s["matrix"]:
            matrix_n = s.get("matrix_n", {})
            survey_n = s["total_respondents"]
            is_oprosso = s.get("data_format") == "oprosso_summary"
            for row_label, answers in s["matrix"].items():
                n_with_experience = matrix_n.get(
                    row_label,
                    _oprosso_matrix_experience_sum(answers),
                )
                sum_row = sum(answers.values())
                excluded_non_usage = sum(
                    cnt
                    for ans, cnt in answers.items()
                    if _category_is_non_usage_only(ans)
                )
                no_answer = max(0, survey_n - sum_row)
                extra = ""
                if is_oprosso:
                    parts = []
                    if excluded_non_usage:
                        parts.append(
                            f"вне базы «с опытом» по колонкам: {excluded_non_usage}"
                        )
                    if no_answer:
                        parts.append(f"не ответили по строке матрицы: {no_answer}")
                    if parts:
                        extra = " (" + "; ".join(parts) + ")"
                    if rh == "off":
                        extra = ""
                else:
                    excluded = survey_n - n_with_experience
                    extra = (
                        f" (исключено {excluded} без опыта)" if excluded else ""
                    )
                    if rh == "off":
                        extra = ""
                lines.append(
                    f"   {row_label}: {n_with_experience} с опытом"
                    + (f" из {survey_n} ответивших на опрос" if is_oprosso else "")
                    + extra
                    + ":"
                )
                exp_items = sorted(
                    [
                        (a, c)
                        for a, c in answers.items()
                        if not _category_is_non_usage_only(a)
                    ],
                    key=lambda x: -x[1],
                )
                non_items = sorted(
                    [
                        (a, c)
                        for a, c in answers.items()
                        if _category_is_non_usage_only(a)
                    ],
                    key=lambda x: -x[1],
                )
                for ans, cnt in exp_items:
                    pct = (
                        100 * cnt / n_with_experience
                        if n_with_experience
                        else 0
                    )
                    lines.append(f"      {ans}: {cnt} ({pct:.1f}%)")
                if is_oprosso and non_items:
                    lines.append(
                        f"      — не в базе «с опытом» (доля от {survey_n} ответивших на опрос):"
                    )
                    for ans, cnt in non_items:
                        pct = 100 * cnt / survey_n if survey_n else 0
                        lines.append(f"      {ans}: {cnt} ({pct:.1f}%)")

        elif s["type"] == "multiple_questions" and s["subquestions"]:
            te_raw = s.get("tool_experience_n")
            is_oprosso_tool = (
                s.get("data_format") == "oprosso_summary" and s.get("tool")
            )
            for sq in s["subquestions"]:
                sq_total = sum(sq["answers"].values())
                if validate:
                    lines.append(f"   {sq['question']} [сумма: {sq_total}]:")
                else:
                    lines.append(f"   {sq['question']}:")
                sq_exp = _tool_page_instrument_answers_adjusted(s, sq["answers"])
                if sq_exp is not None:
                    filt_ans, denom = sq_exp
                    n_sub = _resolve_n_subtract_for_rating_3(s)
                    tr = int(s["total_respondents"] or 0)
                    nu_note = _tool_scale_explanation(s, tr, te_raw, brief, True)
                    if show_bracket:
                        if brief:
                            lines.append(
                                f"      [1–5: из «3» вычтено N={n_sub}; знаменатель %: {denom}. "
                                f"{nu_note}]"
                            )
                        else:
                            lines.append(
                                f"      [Шкала 1–5: из «3» вычитается N={n_sub} ({nu_note}); "
                                f"знаменатель для %: {denom}]"
                            )
                    for ans, cnt in _sort_answers_for_report(list(filt_ans.items())):
                        pct = 100 * cnt / denom if denom else 0
                        lines.append(f"      {ans}: {cnt} ({pct:.1f}%)")
                    continue
                if (
                    show_bracket
                    and is_oprosso_tool
                    and te_raw is not None
                    and sq_total > int(te_raw)
                ):
                    if brief:
                        lines.append(
                            f"      [сумма вариантов {sq_total} > N={te_raw} — % от суммы вариантов]"
                        )
                    else:
                        lines.append(
                            f"      [сумма по вариантам {sq_total}; база «с опытом» N={te_raw} — % от суммы вариантов]"
                        )
                nu_sq = _non_usage_answer_count(sq["answers"])
                if (
                    show_bracket
                    and is_oprosso_tool
                    and te_raw is not None
                    and nu_sq
                ):
                    te_eff = max(1, int(te_raw) - nu_sq)
                    lines.append(
                        f"      [для %: из базы {te_raw} вычтено {nu_sq} по «не использую»/… → {te_eff}]"
                    )
                denom = _denom_oprosso_tool_question(s, sq_total, sq["answers"])
                for ans, cnt in _sort_answers_for_report(list(sq["answers"].items())):
                    pct = 100 * cnt / denom if denom else 0
                    lines.append(f"      {ans}: {cnt} ({pct:.1f}%)")

    return "\n".join(lines)


def build_csv_report(stats_list: List[Dict[str, Any]]) -> str:
    """Формирует CSV-отчёт (плоская структура для импорта в Excel)."""
    import io
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow([
        "Номер", "Страница", "Тип", "Инструмент", "Вопрос/Строка",
        "С опытом (N)", "Вариант ответа", "Кол-во", "%"
    ])

    for s in stats_list:
        n = s["page_num"]
        name = s["page_name"]
        t = s["type"]
        tool = s.get("tool") or ""
        total = s["total_respondents"]

        if t == "single_question":
            total_ans = sum(s["answers"].values())
            exp_only = _tool_page_instrument_answers_adjusted(s, s["answers"])
            if exp_only is not None:
                filt_ans, denom = exp_only
                for ans, cnt in _sort_answers_for_report(list(filt_ans.items())):
                    pct = 100 * cnt / denom if denom else 0
                    w.writerow([n, name, t, tool, "", "", ans, cnt, f"{pct:.1f}"])
            else:
                denom = _denom_oprosso_tool_question(s, total_ans, s["answers"])
                for ans, cnt in _sort_answers_for_report(list(s["answers"].items())):
                    pct = 100 * cnt / denom if denom else 0
                    w.writerow([n, name, t, tool, "", "", ans, cnt, f"{pct:.1f}"])

        elif t == "matrix":
            matrix_n = s.get("matrix_n", {})
            survey_n = s["total_respondents"]
            is_oprosso = s.get("data_format") == "oprosso_summary"
            for row_label, answers in s["matrix"].items():
                n_exp = matrix_n.get(
                    row_label,
                    _oprosso_matrix_experience_sum(answers),
                )
                for ans, cnt in answers.items():
                    if is_oprosso and _category_is_non_usage_only(ans):
                        pct = 100 * cnt / survey_n if survey_n else 0
                        pct_note = f"{pct:.1f} (от {survey_n}, вне базы «с опытом»)"
                    else:
                        pct = 100 * cnt / n_exp if n_exp else 0
                        pct_note = f"{pct:.1f}"
                    w.writerow([n, name, t, tool, row_label, n_exp, ans, cnt, pct_note])

        elif t == "multiple_questions":
            for sq in s["subquestions"]:
                q = sq["question"]
                row_total = sum(sq["answers"].values())
                sq_exp = _tool_page_instrument_answers_adjusted(s, sq["answers"])
                if sq_exp is not None:
                    filt_ans, denom = sq_exp
                    for ans, cnt in _sort_answers_for_report(list(filt_ans.items())):
                        pct = 100 * cnt / denom if denom else 0
                        w.writerow([n, name, t, tool, q, "", ans, cnt, f"{pct:.1f}"])
                else:
                    denom = _denom_oprosso_tool_question(s, row_total, sq["answers"])
                    for ans, cnt in _sort_answers_for_report(list(sq["answers"].items())):
                        pct = 100 * cnt / denom if denom else 0
                        w.writerow([n, name, t, tool, q, "", ans, cnt, f"{pct:.1f}"])

    return buf.getvalue()


def _parse_pages_filter(spec: str) -> Set[int]:
    """Парсит '1,3,5-10' в множество номеров страниц."""
    result = set()
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            try:
                result.update(range(int(a.strip()), int(b.strip()) + 1))
            except ValueError:
                pass
        else:
            try:
                result.add(int(part))
            except ValueError:
                pass
    return result


def build_json_report(stats_list: List[Dict[str, Any]]) -> str:
    """Формирует JSON-отчёт."""
    # Конвертируем defaultdict в dict для JSON
    out = []
    for s in stats_list:
        row = {
            "page_num": s["page_num"],
            "page_name": s["page_name"],
            "type": s["type"],
            "tool": s.get("tool"),
            "total_respondents": s["total_respondents"],
            "answers": s["answers"],
            "matrix": s["matrix"],
            "matrix_n": s.get("matrix_n", {}),
            "subquestions": s["subquestions"],
        }
        if s.get("data_format"):
            row["data_format"] = s["data_format"]
        if s.get("tool_experience_n") is not None:
            row["tool_experience_n"] = s["tool_experience_n"]
        if s.get("tool_without_experience_n") is not None:
            row["tool_without_experience_n"] = s["tool_without_experience_n"]
        if s.get("tool_matrix_non_usage_n") is not None:
            row["tool_matrix_non_usage_n"] = s["tool_matrix_non_usage_n"]
        out.append(row)
    return json.dumps(out, ensure_ascii=False, indent=2)


def analyze_survey_xlsx(filepath: Union[str, Path]) -> List[dict]:
    """
    Анализирует Excel-файл и возвращает список страниц с их типами.

    Возвращает:
        [
            {"row": int, "name": str, "type": str, "description": str},
            ...
        ]
        type: "single_question" | "matrix" | "multiple_questions"
    """
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    result = []
    # Сканируем колонку A: строки вида «число. текст» — границы логических «страниц» вопросника.
    i = 0
    while i < len(rows):
        row = rows[i]
        first = _first_cell(row)

        is_header, page_name, page_num = _is_page_header(first)
        if not is_header or not page_name:
            i += 1
            continue

        next_row = rows[i + 1] if i + 1 < len(rows) else ()
        next_first = _first_cell(next_row)
        next_has_variant = _has_variant_otveta(next_row)

        page_type = "unknown"
        description = ""

        if not next_first or _is_empty_row(next_row):
            page_type = "matrix"
            description = "Вопрос в формате матрицы: по строкам — инструменты, по колонкам — варианты использования"
        elif next_has_variant and VARIANT_OTVETA in next_first:
            page_type = "single_question"
            description = "Страница с одним вопросом и множеством вариантов ответа"
        elif next_has_variant and next_first:
            page_type = "multiple_questions"
            description = "Страница с множеством вопросов (подпункты)"

        result.append({
            "row": i + 1,
            "num": page_num,
            "name": page_name,
            "type": page_type,
            "description": description,
            "tool": extract_tool_from_page_name(page_name),
        })
        i += 1

    return result


def main():
    parser = argparse.ArgumentParser(
        description="Анализ Excel-файла с результатами опроса. "
        "Вывод основных вопросов и фильтрация по ответам на вопрос об инструментах."
    )
    parser.add_argument(
        "file",
        nargs="?",
        default=None,
        help="Путь к xlsx-файлу (по умолчанию result_summary.xlsx или analyze_demo.xlsx)",
    )
    parser.add_argument(
        "--list-questions",
        "-q",
        action="store_true",
        help="Вывести пронумерованный список основных вопросов",
    )
    parser.add_argument(
        "--filter-tool",
        "-f",
        metavar="ИНСТРУМЕНТ",
        help="Применить фильтр: исключить респондентов, ответивших "
        "'не знаю/не использую' на этот инструмент (для страниц детализации)",
    )
    parser.add_argument(
        "--tool-question",
        "-t",
        type=int,
        default=9,
        metavar="N",
        help="Номер вопроса о выборе инструментов (по умолчанию 9)",
    )
    parser.add_argument(
        "--no-filter",
        action="store_true",
        help="Не применять фильтрацию (показать всех респондентов)",
    )
    parser.add_argument(
        "--no-filter-tool-detail",
        action="store_true",
        help="Не исключать по ответам в матрице инструментов "
        "(«не знаком» / «знаю, но не использую») на страницах по инструменту",
    )
    parser.add_argument(
        "--report",
        "-r",
        action="store_true",
        help="Построить расширенный отчёт с числовыми значениями по каждому типу ответов",
    )
    parser.add_argument(
        "--format",
        choices=["text", "csv", "json"],
        default="text",
        help="Формат отчёта (по умолчанию text)",
    )
    parser.add_argument(
        "--output",
        "-o",
        metavar="FILE",
        help="Сохранить отчёт в файл",
    )
    parser.add_argument(
        "--pages",
        "-p",
        metavar="N,M-K",
        help="Только указанные страницы (например: 1,3,5-10)",
    )
    parser.add_argument(
        "--layout",
        choices=["columns", "rows", "auto"],
        default="auto",
        help="Респонденты в столбцах (columns) или в строках (rows). "
        "columns: колонки 1..N = респонденты. rows: строки 1..N = респонденты.",
    )
    parser.add_argument(
        "--validate",
        action="store_true",
        help="Проверка: выводить сумму ответов vs респондентов для каждого вопроса",
    )
    parser.add_argument(
        "--report-hints",
        choices=("full", "brief", "off"),
        default="full",
        metavar="MODE",
        help="Текстовый отчёт: full — подробные пояснения в […]; brief — короче, тот же смысл; "
        "off — только заголовки и числа по вариантам (без пояснений к шкале 1–5 и сводке)",
    )
    parser.add_argument(
        "--exclude-3",
        action="store_true",
        help="Исключать оценку «3» как «нет опыта» (сырой формат и сводка Oprosso); "
        "по умолчанию оценка 3 не исключается",
    )
    parser.add_argument(
        "--exclude-unused-tools",
        action="store_true",
        help="Для oprosso_summary: убрать из отчёта «неиспользуемые» инструменты (см. "
        "--unused-tools-policy) и страницы вопросов по ним",
    )
    parser.add_argument(
        "--unused-tools-policy",
        choices=("regular_or_often", "any_beyond_non_usage"),
        default="regular_or_often",
        help="Критерий «инструмент используется» в матрице вопроса --tool-question: "
        "regular_or_often (по умолчанию) — есть ответы только в «использую регулярно» "
        "или «использую иногда»; any_beyond_non_usage — любой ответ кроме колонок "
        "«не знаком» / «знаю, но не использую» (в т.ч. «пробовал»).",
    )
    args = parser.parse_args()

    filepath = Path(args.file) if args.file else Path("result_summary.xlsx")
    if not filepath.exists():
        filepath = Path("analyze_demo.xlsx")
    if not filepath.exists():
        print(f"Файл не найден: {filepath}", file=sys.stderr)
        sys.exit(1)

    pages = analyze_survey_xlsx(filepath)
    rows = load_survey_rows(filepath)
    total_cols = len(rows[0]) if rows else 0
    data_format = detect_data_format(rows)

    exclusion_map = {}
    if (
        data_format == "raw"
        and not args.no_filter
        and not args.no_filter_tool_detail
    ):
        exclusion_map = build_respondent_exclusion(
            rows, pages, args.tool_question, total_cols
        )

    if args.report:
        layout = args.layout
        if layout == "auto":
            layout = detect_layout(rows)
        stats_list = extract_all_stats(
            rows,
            pages,
            exclusion_map,
            args.tool_question,
            layout,
            exclude_rating_3=args.exclude_3,
            data_format=data_format,
        )
        if args.pages:
            page_nums = _parse_pages_filter(args.pages)
            stats_list = [s for s in stats_list if s["page_num"] in page_nums]
        # Сводка Oprosso: убрать строки/страницы по инструментам без фактического использования
        if args.exclude_unused_tools and data_format == "oprosso_summary":
            unused_keys = compute_unused_tool_keys_oprosso(
                rows,
                pages,
                args.tool_question,
                policy=args.unused_tools_policy,
            )
            print(
                "Политика неиспользуемых инструментов: "
                f"{args.unused_tools_policy} → исключено: {len(unused_keys)}",
                file=sys.stderr,
            )
            if unused_keys:
                print(
                    "Инструменты: "
                    + ", ".join(sorted(unused_keys)),
                    file=sys.stderr,
                )
            stats_list = apply_exclude_unused_tools(stats_list, unused_keys)
        if args.format == "text":
            report_text = build_text_report(
                stats_list,
                validate=args.validate,
                report_hints=args.report_hints,
            )
        elif args.format == "csv":
            report_text = build_csv_report(stats_list)
        else:
            report_text = build_json_report(stats_list)
        if args.output:
            Path(args.output).write_text(report_text, encoding="utf-8")
            print(f"Отчёт сохранён: {args.output}", file=sys.stderr)
        else:
            print(report_text)
        return

    if args.list_questions:
        print("Основные вопросы:\n")
        for num, name in get_main_questions(pages):
            tool = extract_tool_from_page_name(name)
            suffix = f" [инструмент: {tool}]" if tool else ""
            print(f"  {num}. {name}{suffix}")
        print()

    print(f"Файл: {filepath}")
    print(f"Формат данных: {data_format} (oprosso_summary = сводка с «Кол-во ответов»)\n")
    print(f"{'Строка':<6} {'Тип':<20} {'Название страницы'}")
    print("-" * 80)

    for p in pages:
        line = f"{p['row']:<6} {p['type']:<20} {p['name']}"
        filter_info = ""
        if (
            data_format == "raw"
            and not args.no_filter
            and not args.no_filter_tool_detail
            and p.get("tool")
            and total_cols > 1
        ):
            _base = set(range(1, total_cols))
            valid = get_valid_respondent_columns(p["tool"], exclusion_map, _base)
            excluded = len(_base) - len(valid)
            if excluded > 0:
                filter_info = f" (респондентов с фильтром: {len(valid)}/{len(_base)}, исключено: {excluded})"
        elif args.filter_tool and total_cols > 1:
            _base = set(range(1, total_cols))
            valid = get_valid_respondent_columns(
                args.filter_tool, exclusion_map, _base
            )
            excluded = len(_base) - len(valid)
            if p.get("tool") == args.filter_tool and excluded > 0:
                filter_info = f" (учтено: {len(valid)}, исключено: {excluded})"
        print(line + filter_info)
        if p["description"]:
            print(f"       └─ {p['description']}")

    if (
        exclusion_map
        and data_format == "raw"
        and not args.no_filter
        and not args.no_filter_tool_detail
    ):
        print("\n--- Фильтрация страниц по инструменту ---")
        print("Исключаются респонденты с ответами «не знаком» / «знаю, но не использую»")
        print("в матрице инструментов (--tool-question).")
        for tool, cols in sorted(exclusion_map.items()):
            print(f"  {tool}: исключено {len(cols)} респондентов")


if __name__ == "__main__":
    main()
